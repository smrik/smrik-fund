"""Excel forecast edits compiled against stable row anchors; sources stay generated.

Create once: python -m smrik_fund.workbook_template init <reviewed/run/directory>
Validate edits: python -m smrik_fund.workbook_template check
"""

import argparse
import hashlib
import json
import math
import re
from io import BytesIO
from pathlib import Path
from zipfile import ZipFile

from openpyxl import load_workbook
from openpyxl.workbook.defined_name import DefinedName

ROOT = Path(__file__).resolve().parents[2]
DEFAULT_TEMPLATE = ROOT / "templates/company-model.xlsx"
EDITABLE = ("Income", "BalanceSheet", "CashFlow", "Assets", "WorkingCapital", "Stub")
REF = re.compile(
	r"(?<![A-Za-z0-9_])(?:(?:'([^']+)'|([A-Za-z_][A-Za-z0-9_]*))!)?(\$?[A-Z]{1,3})(\$?)([1-9][0-9]*)(?![A-Za-z0-9_(])"
)


def _read_book(path):
	with ZipFile(path) as archive:
		if any(
			"externalLink" in n or "vbaProject" in n or "connections" in n
			for n in archive.namelist()
		):
			raise ValueError(
				"Template must contain local formulas only, without external links or macros"
			)
	book = load_workbook(path, data_only=False)
	try:
		sheets = {}
		for sheet in book:
			if sheet.max_row > 1000 or sheet.max_column > 100:
				raise ValueError("Template exceeds the supported model area")
			cells = {}
			for row in sheet:
				for cell in row:
					value = cell.value
					if value is None or value == "":
						continue
					if cell.data_type == "f":
						if not isinstance(value, str):
							raise ValueError(
								"Use ordinary cell formulas, not arrays or data tables"
							)
						value = {"formula": value}
					cells[cell.coordinate] = value
			sheets[sheet.title] = cells
		return sheets, {key: name.attr_text for key, name in book.defined_names.items()}
	finally:
		book.close()


def _move_formula(formula, sheet, maps):
	# Ignore references embedded in Excel string literals.
	parts = re.split(r'("(?:[^"]|"")*")', formula)
	for i in range(0, len(parts), 2):
		last_end, last_sheet = -1, sheet

		def replace(match, part=parts[i]):
			nonlocal last_end, last_sheet
			s = (
				match[1]
				or match[2]
				or (
					last_sheet
					if part[last_end : match.start()].strip() == ":"
					else sheet
				)
			)
			last_end, last_sheet = match.end(), s
			r = maps.get(s, {}).get(match[5], int(match[5]))
			prefix = f"{s}!" if match[1] or match[2] else ""
			return f"{prefix}{match[3]}{match[4]}{r}"

		parts[i] = REF.sub(replace, parts[i])
	return "".join(parts)


def initialize(reviewed_dir, target=DEFAULT_TEMPLATE):
	"""Bind a known workbook and snapshot, never overwrite a user's template."""
	reviewed_dir, target = Path(reviewed_dir), Path(target)
	contract = target.with_suffix(".json")
	if target.exists() or contract.exists():
		raise ValueError(
			"Template already exists; choose another path to preserve your edits"
		)
	model = json.loads((reviewed_dir / "model.json").read_text())
	snapshot = json.loads((reviewed_dir / "snapshot.json").read_text())
	if snapshot["mechanical"] != "PASS" or snapshot["valuation_gate"] != "PASS":
		raise ValueError(
			"Initialize from a model that passed accounting and valuation checks"
		)
	previous_patch = model.get("workbook_template", {})
	if previous_patch.get("changes") or any(
		int(r) != moved
		for rows in previous_patch.get("row_maps", {}).values()
		for r, moved in rows.items()
	):
		raise ValueError(
			"Initialize from an uncustomized generated workbook; copy an existing template to preserve its edits"
		)
	source = reviewed_dir / f"{model['case']}.xlsx"
	sheets, _ = _read_book(source)
	baseline = {
		"schema": "company-excel-template-v1",
		"example_ticker": model["case"],
		"history_columns": snapshot["presentation"]["history_columns"],
		"forecast_columns": len(snapshot["presentation"]["headers"])
		- snapshot["presentation"]["history_columns"]
		- 2,
		"input_rows": snapshot["input_rows"],
		"sheets": sheets,
	}
	target.parent.mkdir(parents=True, exist_ok=True)
	book = load_workbook(source)
	try:
		for name in EDITABLE:
			last = max(int(re.search(r"\d+$", c)[0]) for c in sheets[name])
			for row in range(1, last + 1):
				book.defined_names.add(
					DefinedName(f"_smrik_{name}_{row}", attr_text=f"'{name}'!$A${row}")
				)
		book.save(target)
	finally:
		book.close()
	baseline["sheets"] = _read_book(target)[0]
	contract.write_text(json.dumps(baseline, indent=2) + "\n", encoding="utf-8")
	return target


def compile_template(path=DEFAULT_TEMPLATE):
	"""Return changed cells and row maps, not copied company financial data."""
	path = Path(path)
	if path.suffix.lower() != ".xlsx":
		raise ValueError("Template must be an .xlsx workbook")
	contract_path = path.with_suffix(".json")
	contract_bytes, template_bytes = contract_path.read_bytes(), path.read_bytes()
	baseline = json.loads(contract_bytes)
	if baseline.get("schema") != "company-excel-template-v1":
		raise ValueError("Unsupported template contract")
	sheets, names = _read_book(BytesIO(template_bytes))
	if set(sheets) != set(baseline["sheets"]):
		raise ValueError(
			"Keep the existing sheet names; extra sheets are not supported"
		)
	maps = {}
	for sheet in EDITABLE:
		last = max(int(re.search(r"\d+$", c)[0]) for c in baseline["sheets"][sheet])
		mapping = {}
		for row in range(1, last + 1):
			anchor = names.get(f"_smrik_{sheet}_{row}", "")
			match = re.fullmatch(rf"'?{sheet}'?!\$A\$(\d+)", anchor)
			if not match:
				raise ValueError(
					f"Missing or broken row anchor: {sheet} original row {row}"
				)
			mapping[str(row)] = int(match[1])
		if list(mapping.values()) != sorted(set(mapping.values())) or any(
			mapping[str(r)] != r for r in (1, 2, 3)
		):
			raise ValueError(
				"Keep original rows in order and the first three header rows in place"
			)
		maps[sheet] = mapping
	changes = {}
	first_forecast = baseline["history_columns"] + 3  # one-based column
	for sheet, cells in baseline["sheets"].items():
		expected = {}
		for address, value in cells.items():
			column, row = re.fullmatch(r"([A-Z]+)(\d+)", address).groups()
			new_address = column + str(maps.get(sheet, {}).get(row, int(row)))
			if isinstance(value, dict):
				value = {"formula": _move_formula(value["formula"], sheet, maps)}
			expected[new_address] = value
		for address in sorted(set(expected) | set(sheets[sheet])):
			value = sheets[sheet].get(address)
			if isinstance(value, dict):
				value = {"formula": _move_formula(value["formula"], sheet, {})}
			previous = expected.get(address)
			if value == previous or (
				isinstance(value, (int, float))
				and isinstance(previous, (int, float))
				and math.isclose(value, previous, rel_tol=1e-14, abs_tol=1e-12)
			):
				continue
			column, row = re.fullmatch(r"([A-Z]+)(\d+)", address).groups()
			c = 0
			for char in column:
				c = c * 26 + ord(char) - 64
			editable = (
				sheet in EDITABLE
				and int(row) > 3
				and (
					c == 1
					or (sheet == "Stub" and c == 2)
					or (
						sheet != "Stub"
						and first_forecast
						<= c
						< first_forecast + baseline.get("forecast_columns", 10)
					)
				)
			)
			if not editable:
				raise ValueError(
					f"Protected source, header or audit cell changed: {sheet}!{address}"
				)
			if isinstance(value, dict):
				formula = value["formula"]
				if re.search(
					r"[\[\]#|]|\b(INDIRECT|OFFSET|HYPERLINK|WEBSERVICE|RTD|DDE)\s*\(",
					formula,
					re.I,
				):
					raise ValueError(f"Unsupported template formula: {sheet}!{address}")
				for ref in REF.finditer(formula):
					if (ref[1] or ref[2] or sheet) not in (*EDITABLE, "Inputs"):
						raise ValueError(
							"Edited formulas may reference model tabs and Inputs only"
						)
			changes.setdefault(sheet, {})[address] = value
	book = load_workbook(BytesIO(template_bytes), data_only=False)
	try:
		formats = {
			sheet: {cell: book[sheet][cell].number_format for cell in cells}
			for sheet, cells in changes.items()
		}
	finally:
		book.close()
	return {
		"schema": baseline["schema"],
		"sha256": hashlib.sha256(template_bytes).hexdigest(),
		"contract_sha256": hashlib.sha256(contract_bytes).hexdigest(),
		"example_ticker": baseline["example_ticker"],
		"history_columns": baseline["history_columns"],
		"forecast_columns": baseline.get("forecast_columns", 10),
		"input_rows": baseline["input_rows"],
		"row_maps": maps,
		"changes": changes,
		"number_formats": formats,
		"note": "Only forecast edits are reused. Fresh company sources and audit checks are generated each run. Custom formulas require financial review.",
	}


def main():
	parser = argparse.ArgumentParser(description=__doc__)
	parser.add_argument("command", choices=("init", "check"))
	parser.add_argument("reviewed_dir", nargs="?", type=Path)
	parser.add_argument("--template", type=Path, default=DEFAULT_TEMPLATE)
	args = parser.parse_args()
	if args.command == "init":
		if args.reviewed_dir is None:
			parser.error("init requires a reviewed run directory")
		initialize(args.reviewed_dir, args.template)
	patch = compile_template(args.template)
	print(
		json.dumps(
			{
				"template": str(args.template),
				"sha256": patch["sha256"],
				"edited_cells": sum(len(x) for x in patch["changes"].values()),
				"status": "VALID",
			},
			indent=2,
		)
	)


if __name__ == "__main__":
	main()
