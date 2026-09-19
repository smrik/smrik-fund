import copy
import json
import shutil
import subprocess
from pathlib import Path

import pytest
from openpyxl import load_workbook

from smrik_fund.workbook_template import DEFAULT_TEMPLATE, compile_template

ROOT = Path(__file__).resolve().parents[1]


def template_copy(tmp_path):
	path = tmp_path / "custom.xlsx"
	shutil.copyfile(DEFAULT_TEMPLATE, path)
	shutil.copyfile(DEFAULT_TEMPLATE.with_suffix(".json"), path.with_suffix(".json"))
	return path


def test_default_template_is_an_identity_and_has_source_bindings():
	patch = compile_template()
	assert patch["changes"] == {}
	assert patch["history_columns"] == 5
	assert patch["input_rows"]["beta"] > 0
	assert patch["row_maps"]["Income"]["38"] == 38


def test_native_excel_insertions_preserve_anchors_and_only_capture_actual_edits(
	tmp_path,
):
	path = template_copy(tmp_path)
	shutil.copyfile(ROOT / "tests/fixtures/templates/inserted-expense.xlsx", path)
	patch = compile_template(path)
	assert patch["row_maps"]["Income"]["38"] == 40
	assert sum(len(cells) for cells in patch["changes"].values()) == 32
	assert patch["changes"]["Income"]["H32"] == {"formula": "=H8*H33"}
	assert patch["number_formats"]["Income"]["H33"] == "0.0%"


@pytest.mark.parametrize(
	"sheet,cell,value",
	[
		("Income", "B8", 999999),
		("Inputs", "B4", 999999),
		("Checks", "B9", "PASS"),
		("DCF", "B19", 999),
		("Income", "A3", "wrong periods"),
	],
)
def test_reported_values_periods_and_checks_cannot_be_reused_as_edits(
	tmp_path, sheet, cell, value
):
	path = template_copy(tmp_path)
	book = load_workbook(path)
	book[sheet][cell] = value
	book.save(path)
	book.close()
	with pytest.raises(ValueError, match="Protected source, header or audit cell"):
		compile_template(path)


def test_excel_formula_edit_is_preserved_without_cached_values(tmp_path):
	path = template_copy(tmp_path)
	book = load_workbook(path)
	book["Income"]["H32"] = "=H8-H26+H29-100"
	book.save(path)
	book.close()
	patch = compile_template(path)
	assert patch["changes"] == {"Income": {"H32": {"formula": "=H8-H26+H29-100"}}}
	assert patch["sha256"] != compile_template()["sha256"]


def test_deleted_row_anchor_stops_before_calculation(tmp_path):
	path = template_copy(tmp_path)
	book = load_workbook(path)
	book.defined_names["_smrik_Income_38"].attr_text = "#REF!"
	book.save(path)
	book.close()
	with pytest.raises(ValueError, match="Missing or broken row anchor"):
		compile_template(path)


@pytest.mark.parametrize(
	"formula", ['=INDIRECT("H8")', "='[other.xlsx]Income'!H8", "=History!B4"]
)
def test_unsupported_formula_dependencies_stop(tmp_path, formula):
	path = template_copy(tmp_path)
	book = load_workbook(path)
	book["Income"]["H32"] = formula
	book.save(path)
	book.close()
	with pytest.raises(
		ValueError, match="Unsupported template formula|Edited formulas may reference"
	):
		compile_template(path)


def inserted_cost_patch():
	patch = compile_template()
	patch["row_maps"]["Income"] = {
		r: int(r) + (2 if int(r) >= 32 else 0) for r in patch["row_maps"]["Income"]
	}
	patch["changes"] = {
		"Income": {
			"A32": "Additional scenario operating expense",
			"A33": "Additional expense / revenue",
		}
	}
	for column in "HIJKLMNOPQ":
		patch["changes"]["Income"].update(
			{
				f"{column}32": {"formula": f"={column}8*{column}33"},
				f"{column}33": 0.005,
				f"{column}34": {
					"formula": f"={column}8-{column}26+{column}29-{column}32"
				},
			}
		)
	return patch


def test_cross_sheet_ranges_absolute_refs_input_bindings_and_history_columns(tmp_path):
	patch = inserted_cost_patch()
	patch["changes"]["Income"]["H33"] = {
		"formula": f"=Inputs!$B${patch['input_rows']['cogs_ratio']}"
	}
	payload = tmp_path / "patch.json"
	payload.write_text(json.dumps(patch))
	script = """
import {readFileSync} from 'node:fs';
import {templateEdits} from './scripts/spreadsheet_compat/company-template.mjs';
const patch=JSON.parse(readFileSync(process.argv[1]));
const e=templateEdits(patch,3,{cogs_ratio:80});
const g=e.table('Income',Array.from({length:57},()=>[]));
console.log(JSON.stringify({
 range:e.formula('=SUM(Income!$H$32:$H$38)', 'CashFlow'),
 text:e.formula('=IF(A1="Income!H38",Income!H38,0)', 'CashFlow'),
 custom:g[32][5].formula,
 expense:g[31][5].formula,
 binding:e.address({sheet:'Income',cell:'F38'}),
 tail:e.address({sheet:'Income',cell:'F62'}),
 horizon:(()=>{try{templateEdits(patch,3,{},11);return false;}catch{return true;}})()
}));
"""
	result = subprocess.run(
		["node", "--input-type=module", "-e", script, str(payload)],
		cwd=ROOT,
		check=True,
		text=True,
		capture_output=True,
	)
	assert json.loads(result.stdout) == {
		"range": "=SUM(Income!$H$34:$H$40)",
		"text": '=IF(A1="Income!H38",Income!H40,0)',
		"custom": "=Inputs!$B$80",
		"expense": "=F8*F33",
		"binding": {"sheet": "Income", "cell": "F40"},
		"tail": {"sheet": "Income", "cell": "F64"},
		"horizon": True,
	}


@pytest.mark.parametrize(
	"ticker,case",
	[
		("BBWI", "data/history/20260919/BBWI/source"),
		("AAPL", "data/second-company/AAPL/frozen-20260910"),
		("NVDA", "data/multi-ticker/20260910/NVDA/attempt-r2/source"),
	],
)
def test_inserted_expense_recalculates_all_statements_for_different_tickers(
	tmp_path, ticker, case
):
	from smrik_fund.company_run import build
	from smrik_fund.portable_model import prepare_model

	if not (ROOT / case).exists():
		pytest.skip("Local frozen acceptance case required")
	model = prepare_model(ROOT / case)
	base = build(model, tmp_path / "base")
	custom = copy.deepcopy(model)
	custom["workbook_template"] = inserted_cost_patch()
	changed = build(custom, tmp_path / "edited")
	assert changed["per_share_value"] < base["per_share_value"]
	assert changed["schedules"]["Revenue"] == base["schedules"]["Revenue"]
	for i in range(1, 11):
		assert base["schedules"]["EBIT"][i] - changed["schedules"]["EBIT"][
			i
		] == pytest.approx(base["schedules"]["Revenue"][i] * 0.005)
		assert (
			abs(changed["schedules"]["Balance difference (no cash/equity plug)"][i])
			< 1e-6
		)
	assert all(changed["proof"].values())
	assert len(changed["template_cells"]) == 32


def test_unbalanced_template_cannot_publish_a_valuation(tmp_path):
	from smrik_fund.company_run import build
	from smrik_fund.portable_model import prepare_model

	case = ROOT / "data/history/20260919/BBWI/source"
	if not case.exists():
		pytest.skip("Local frozen acceptance case required")
	model = prepare_model(case)
	patch = compile_template()
	patch["changes"] = {"BalanceSheet": {"H15": {"formula": "=1"}}}
	model["workbook_template"] = patch
	with pytest.raises(ValueError, match="Calculated company model is blocked"):
		build(model, tmp_path / "bad")
	assert not (tmp_path / "bad/version.json").exists()


def test_analyst_and_reviewer_changes_recalculate_before_acceptance_without_api(
	tmp_path, monkeypatch
):
	"""Exercise orchestration with simulated responses; this is not a live LLM test."""
	from smrik_fund import company_run

	case = ROOT / "data/history/20260919/BBWI/source"
	if not case.exists():
		pytest.skip("Local frozen acceptance case required")
	prices = tmp_path / "prices"
	prices.mkdir()
	for name in ("luna", "sol"):
		(prices / f"{name}-price-snapshot.json").write_text(
			'{"model":"simulated-test"}'
		)
	calls = []

	def simulated_call(output_dir, name, payload, **kwargs):
		calls.append((name, copy.deepcopy(payload)))
		controls = dict(payload["model"]["controls"])
		answer = {
			"controls": controls,
			"rationale": "SIMULATED TEST ONLY",
			"limitations": [],
		}
		if name == "analyst":
			controls["products_growth"] = -0.02
			controls["payout_ratio"] = 0
		elif name == "review-0":
			controls["cogs_ratio"] += 0.01
			answer["verdict"] = "revise"
		else:
			answer["verdict"] = "accept"
		return answer

	monkeypatch.setattr(company_run, "call_model", simulated_call)
	out = tmp_path / "simulated-run"
	result = company_run.run_case(
		"BBWI", case, out, live=True, price_dir=prices, native_excel=False
	)
	assert [name for name, _ in calls] == ["analyst", "review-0", "review-1"]
	first, revised = [
		payload["calculated"] for name, payload in calls if name.startswith("review")
	]
	assert revised["per_share_value"] < first["per_share_value"]
	assert revised["schedules"]["Revenue"] == first["schedules"]["Revenue"]
	assert result["per_share_value"] == revised["per_share_value"]
	assert calls[1][1]["model"]["controls"]["products_growth"] == -0.02
	assert (out / "template.xlsx").is_file()
	assert result["template_sha256"] == compile_template()["sha256"]
	assert not (out / "analyst.receipt.json").exists()  # No paid transport invoked.
